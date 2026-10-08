"""A model's parameters as one spreadsheet and back (STD-36): the examples
round-trip with no differences, edits come back as a list of changes, and
rows that do not fit the model are refused with their row numbers."""
from __future__ import annotations

import base64
import io
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.dataio import params
from app.dataio.cli import main as cli_main
from app.main import app
from app.storage import load_example

client = TestClient(app)
EXAMPLES = ["bev-car", "hybrid-car", "fs-electric"]


@pytest.mark.parametrize("example", EXAMPLES)
@pytest.mark.parametrize("kind", ["xlsx", "csv"])
def test_the_examples_round_trip_with_no_differences(example, kind):
    project = load_example(example)
    data = params.export_xlsx(project) if kind == "xlsx" else params.export_csv(project)
    res = params.import_sheet(project, data, f"p.{kind}")
    assert res.errors == []
    assert res.changes == []
    assert res.rows == res.unchanged > 50


def _edit(project, edits: dict[tuple[str, str], object], table_edit=None) -> bytes:
    """The project's sheet with Value cells changed (by part id and key)."""
    wb = openpyxl.load_workbook(io.BytesIO(params.export_xlsx(project)))
    ws = wb["Parameters"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    for row in range(2, ws.max_row + 1):
        key = (ws.cell(row, col["Part ID"]).value, ws.cell(row, col["Key"]).value)
        for (pid, k), change in edits.items():
            if key == (pid, k):
                field, value = change if isinstance(change, tuple) else ("Value", change)
                ws.cell(row, col[field]).value = value
    if table_edit:
        table_edit(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row_of(project, pid: str, key: str) -> int:
    rows = params.export_sheets(project)[0].rows
    return next(i + 1 for i, r in enumerate(rows) if r[1] == pid and r[4] == key)


def test_edits_come_back_as_changes():
    project = load_example("bev-car")
    vehicle = next(e for e in project.systems[0].elements if e.componentDefId == "vehicle.body")
    diff = next(e for e in project.systems[0].elements if e.componentDefId == "mech.differential")

    def ocv(wb):
        ws = wb["HV Battery Pack - Open-Circuit"]
        ws["B3"] = 333.5  # the first OCV point

    data = _edit(project, {(vehicle.id, "mass_kg"): 2300, (diff.id, "locked"): "TRUE"}, ocv)
    res = params.import_sheet(project, data, "p.xlsx")
    assert res.errors == []
    got = {(c.elementId, c.key): c for c in res.changes}
    assert got[(vehicle.id, "mass_kg")].new == 2300
    assert got[(vehicle.id, "mass_kg")].row == _row_of(project, vehicle.id, "mass_kg")
    assert got[(diff.id, "locked")].new is True
    battery = next(e for e in project.systems[0].elements if e.componentDefId == "battery.generic")
    ocv_change = got[(battery.id, "ocv_table")]
    assert ocv_change.new["0"] == 333.5 and ocv_change.sheet == "HV Battery Pack - Open-Circuit"
    assert len(res.changes) == 3


def test_five_unit_errors_give_five_row_numbered_messages():
    project = load_example("fs-electric")
    els = project.systems[0].elements
    vehicle = next(e for e in els if e.componentDefId == "vehicle.body")
    motor = next(e for e in els if e.componentDefId == "motor.emotor")
    battery = next(e for e in els if e.componentDefId == "battery.generic")
    seeded = {
        (vehicle.id, "mass_kg"): ("Unit", "lb"),
        (vehicle.id, "frontal_area_m2"): ("Unit", "ft²"),
        (motor.id, "max_speed_rpm"): ("Unit", "rad/s"),
        (battery.id, "capacity_kWh"): ("Unit", "MJ"),
        (battery.id, "max_charge_power_kW"): ("Unit", "hp"),
    }
    res = params.import_sheet(project, _edit(project, seeded), "p.xlsx")
    assert len(res.errors) == 5
    rows = sorted(_row_of(project, pid, key) for pid, key in seeded)
    assert sorted(e["row"] for e in res.errors) == rows
    for e in res.errors:
        assert e["text"].startswith(f"Row {e['row']}: ") and "Write the value in" in e["text"]


def test_rows_that_do_not_fit_the_model_are_refused():
    project = load_example("bev-car")
    vehicle = next(e for e in project.systems[0].elements if e.componentDefId == "vehicle.body")
    wheel = next(e for e in project.systems[0].elements if e.componentDefId == "propulsion.wheel")
    data = _edit(project, {
        (vehicle.id, "mass_kg"): "heavy",
        (wheel.id, "axle"): "Middle",
        (vehicle.id, "frontal_area_m2"): ("Key", "no_such_key"),
        (wheel.id, "radius_m"): ("Part ID", "el-gone"),
    })
    texts = " | ".join(e["text"] for e in params.import_sheet(project, data, "p.xlsx").errors)
    assert "'heavy' is not a number" in texts
    assert "'Middle' is not one of: Front, Rear" in texts
    assert "has no parameter 'no_such_key'" in texts
    assert "no part with id 'el-gone'" in texts


def test_a_value_outside_its_limits_is_a_warning():
    project = load_example("bev-car")
    vehicle = next(e for e in project.systems[0].elements if e.componentDefId == "vehicle.body")
    res = params.import_sheet(project, _edit(project, {(vehicle.id, "mass_kg"): -5}), "p.xlsx")
    assert res.errors == [] and len(res.changes) == 1
    assert "Data Checks will report it" in res.warnings[0]["text"]


def test_a_sheet_from_german_excel_reads_dots_between_thousands():
    project = load_example("bev-car")
    vehicle = next(e for e in project.systems[0].elements if e.componentDefId == "vehicle.body")
    data = (f"Part ID;Key;Value\n{vehicle.id};mass_kg;1.650\n"
            "el-wheel-rl;radius_m;0,35\n").encode("cp1252")
    res = params.import_sheet(project, data, "p.csv")
    assert res.errors == [] and res.warnings == []
    assert {(c.key, c.new) for c in res.changes} == {("mass_kg", 1650), ("radius_m", 0.35)}
    # 1.650 alone could be 1.65: the sheet says how it was read
    res = params.import_sheet(project, f"Part ID;Key;Value\n{vehicle.id};mass_kg;1.650\n"
                              .encode(), "p.csv")
    assert res.changes[0].new == 1650 and "save the sheet as .xlsx" in res.warnings[0]["text"]


def test_a_sheet_without_the_columns_is_refused():
    res = params.import_sheet(load_example("bev-car"), b"a,b\n1,2\n", "x.csv")
    assert "Part ID, Key and Value" in res.errors[0]["text"]


def test_the_api_exports_imports_and_offers_the_template():
    project = load_example("bev-car").model_dump(mode="json")
    r = client.post("/api/params/export", json={"project": project, "format": "xlsx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "parameters.xlsx" in r.headers["content-disposition"]
    back = client.post("/api/params/import", json={
        "project": project, "filename": "p.xlsx",
        "data": base64.b64encode(r.content).decode()}).json()
    assert back["ok"] and back["changes"] == [] and back["unchanged"] == back["rows"]
    r = client.post("/api/params/export", json={"project": project, "format": "csv"})
    assert r.content.startswith(b"\xef\xbb\xbf")
    t = client.get("/api/params/template")
    assert t.status_code == 200 and "Formula Student" in t.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(t.content))
    assert wb.sheetnames[0] == "Parameters" and "About" in wb.sheetnames


def test_params_on_the_command_line(tmp_path, capsys):
    project = load_example("bev-car")
    path = tmp_path / "bev.json"
    path.write_text(project.model_dump_json(), encoding="utf-8")
    sheet = tmp_path / "p.xlsx"
    assert cli_main(["params", "export", str(path), "--out", str(sheet)]) == 0
    vehicle = next(e for e in project.systems[0].elements if e.componentDefId == "vehicle.body")
    sheet.write_bytes(_edit(project, {(vehicle.id, "mass_kg"): 2100}))
    new = tmp_path / "bev-2100.json"
    assert cli_main(["params", "import", str(path), str(sheet), "--out", str(new)]) == 0
    assert "1 change(s)" in capsys.readouterr().out
    saved = json.loads(new.read_text("utf-8"))
    el = next(e for s in saved["systems"] for e in s["elements"] if e["id"] == vehicle.id)
    assert el["parameterOverrides"]["mass_kg"] == 2100
